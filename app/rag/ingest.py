"""
知识库文档入库脚本

把 docs/knowledge_base/ 下的文档构建成本地 FAISS 向量集合：

    docs/knowledge_base/
    ├── 电商行业/          -> 集合「电商行业」
    │   ├── xxx.pdf
    │   └── yyy.docx
    └── 金融行业/          -> 集合「金融行业」
        └── zzz.pdf

每个一级子目录视为一个独立集合；集合内的 PDF/DOCX/MD/TXT 会被解析、
分片并向量化，结果写入 app/data/vector_store/<集合名>/。

用法：
    uv run python -m app.rag.ingest                      # 使用默认路径
    uv run python -m app.rag.ingest --source docs/knowledge_base --store app/data/vector_store
    uv run python -m app.rag.ingest --collection 金融行业  # 只重建指定集合
"""

import argparse
import datetime
from pathlib import Path
from typing import Optional

from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.rag.embeddings import get_embeddings
from app.rag.vector_store import VectorKnowledgeBase
from app.utils.logging_config import get_logger, setup_logging

logger = get_logger(__name__)

# 中文报告类文档：块大小 500 字符左右，重叠 80 字符，保留段落边界
CHUNK_SIZE = 500
CHUNK_OVERLAP = 80

# 当前文件位于 app/rag/ingest.py，parents[1] 即 app 目录
APP_ROOT = Path(__file__).parents[1]
PROJECT_ROOT = APP_ROOT.parent

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".md", ".txt"}


def load_single_file(file_path: Path) -> list[Document]:
    """
    解析单个文件为 LangChain Document 列表

    支持 PDF（pypdf）、Word（python-docx）、Markdown 和纯文本。
    :param file_path: 文件路径
    :return: Document 列表；解析失败时返回空列表并记录警告
    """
    suffix = file_path.suffix.lower()
    try:
        if suffix == ".pdf":
            import pypdf

            reader = pypdf.PdfReader(str(file_path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            import docx

            document = docx.Document(str(file_path))
            text = "\n".join(para.text for para in document.paragraphs)
        elif suffix in {".md", ".txt"}:
            text = file_path.read_text(encoding="utf-8")
        else:
            logger.warning("跳过不支持的文件类型: %s", file_path.name)
            return []
    except Exception as e:
        logger.warning("文件解析失败，已跳过: %s (%s)", file_path.name, e)
        return []

    if not text.strip():
        logger.warning("文件无可提取文本，已跳过: %s", file_path.name)
        return []

    # source 元数据让检索结果可以标注出处文件名
    return [Document(page_content=text, metadata={"source": file_path.name})]


def load_collection_documents(source_dir: Path) -> list[Document]:
    """
    递归加载一个集合目录下的全部文档

    :param source_dir: 集合源目录
    :return: Document 列表
    """
    documents: list[Document] = []
    for file_path in sorted(source_dir.rglob("*")):
        if file_path.is_file() and file_path.suffix.lower() in SUPPORTED_SUFFIXES:
            documents.extend(load_single_file(file_path))
    return documents


def split_documents(documents: list[Document]) -> list[Document]:
    """
    对文档做递归字符分片

    :param documents: 原始文档列表
    :return: 分片后的文档列表
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        # 中文场景优先按段落、换行和标点切分，尽量保持语义完整
        separators=["\n\n", "\n", "。", "；", "！", "？", "，", " ", ""],
    )
    return splitter.split_documents(documents)


def build_collection(
    name: str,
    source_dir: Path,
    knowledge_base: VectorKnowledgeBase,
) -> Optional[int]:
    """
    构建单个知识库集合

    :param name: 集合名称
    :param source_dir: 集合源文档目录
    :param knowledge_base: 向量库管理器
    :return: 成功时返回分片数量；源目录无有效文档时返回 None
    """
    documents = load_collection_documents(source_dir)
    if not documents:
        logger.warning("集合 '%s' 源目录无有效文档，跳过: %s", name, source_dir)
        return None

    chunks = split_documents(documents)
    file_names = sorted({doc.metadata.get("source", "") for doc in documents})

    logger.info(
        "开始构建集合 '%s': %d 个文件 -> %d 个分片",
        name,
        len(documents),
        len(chunks),
    )
    store = FAISS.from_documents(chunks, get_embeddings())
    embedding_model = getattr(get_embeddings(), "model", "")
    knowledge_base.save_collection(
        name=name,
        store=store,
        chunk_count=len(chunks),
        file_names=file_names,
        embedding_model=embedding_model,
    )
    return len(chunks)


def main() -> None:
    """命令行入口：扫描源目录并构建全部（或指定）集合"""
    parser = argparse.ArgumentParser(description="构建本地 FAISS 向量知识库")
    parser.add_argument(
        "--source",
        type=Path,
        default=PROJECT_ROOT / "docs" / "knowledge_base",
        help="知识库源文档根目录（每个一级子目录为一个集合）",
    )
    parser.add_argument(
        "--store",
        type=Path,
        default=APP_ROOT / "data" / "vector_store",
        help="向量库输出目录",
    )
    parser.add_argument(
        "--collection",
        type=str,
        default=None,
        help="只构建指定名称的集合；默认构建全部",
    )
    args = parser.parse_args()

    setup_logging("INFO")

    if not args.source.exists():
        logger.error("源目录不存在: %s", args.source)
        return

    knowledge_base = VectorKnowledgeBase(args.store)

    collection_dirs = [d for d in sorted(args.source.iterdir()) if d.is_dir()]
    if args.collection:
        collection_dirs = [d for d in collection_dirs if d.name == args.collection]
        if not collection_dirs:
            logger.error("未找到指定集合目录: %s", args.collection)
            return

    started = datetime.datetime.now()
    built = 0
    for collection_dir in collection_dirs:
        if build_collection(collection_dir.name, collection_dir, knowledge_base):
            built += 1

    logger.info(
        "入库完成: %d/%d 个集合，耗时 %.1fs",
        built,
        len(collection_dirs),
        (datetime.datetime.now() - started).total_seconds(),
    )


if __name__ == "__main__":
    main()
